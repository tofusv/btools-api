"""Offline safety/integrity tests. Run directly; never call the Gemini API."""
from pathlib import Path
import base64
import copy
import importlib
from io import BytesIO
import json
import os
import sys
import tempfile
import time
import unittest
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / '.test_dependencies'))
sys.path.insert(0, str(ROOT / 'btools_webapp'))
from docx import Document
from docx.oxml.ns import qn
from fastapi.testclient import TestClient
from job_store import JobStore, JobProcessingError
from reference_parser import source_from_text, resolve_plan, validate_source, build_prompt
from source_reader import source_from_docx_bytes
from generate_course_outline import generate_doc
api = importlib.import_module('app')


def node(nid, kind, bid=None, text='', parent='', start=0, end=None, table='', category=''):
    return {'id': nid, 'parent_id': parent, 'kind': kind, 'category': category,
            'refs': [{'block_id': bid, 'start': start, 'end': len(text) if end is None else end}] if bid else [],
            'table_id': table, 'needs_review': False}


def plan(nodes, edits=None):
    return {'nodes': nodes, 'edits': edits or [], 'generated_title_en': '', 'warnings': []}


def doc_bytes(doc):
    stream = BytesIO(); doc.save(stream); return stream.getvalue()


def identity_plan(source):
    """Fixture for integrity checks only; this is not an AI quality benchmark."""
    table_for_block = {bid: t['id'] for t in source['tables'] for r in t['rows'] for c in r['cells'] for bid in c['block_ids']}
    nodes, emitted = [], set()
    for b in source['blocks']:
        tid = table_for_block.get(b['id'])
        if tid:
            if tid not in emitted:
                nodes.append(node('n' + str(len(nodes)), 'table', table=tid)); emitted.add(tid)
        else:
            nodes.append(node('n' + str(len(nodes)), 'paragraph', b['id'], b['text']))
    return plan(nodes)


class ReferenceTests(unittest.TestCase):
    def test_nested_modules_and_source_table_render_without_losing_content(self):
        original = Document()
        for text in ['Day 1', '09:00', 'Module A', 'Exercise', 'Day 2', 'Module B', 'Conclusion']:
            original.add_paragraph(text)
        original.add_table(rows=1, cols=1).cell(0, 0).text = 'Original table'
        source = source_from_docx_bytes(doc_bytes(original))
        specs = [('day1', 'module', ''), ('time', 'time', 'day1'),
                 ('a', 'module', 'day1'), ('exercise', 'paragraph', 'a'),
                 ('day2', 'module', ''), ('b', 'module', 'day2'),
                 ('conclusion', 'paragraph', 'b')]
        nodes = [node(nid, kind, block['id'], block['text'], parent)
                 for (nid, kind, parent), block in zip(specs, source['blocks'])]
        nodes.append(node('table', 'table', parent='b', table=source['tables'][0]['id']))
        with tempfile.TemporaryDirectory(dir=ROOT) as tmp:
            path = Path(tmp) / 'nested.docx'
            generate_doc(resolve_plan(source, plan(nodes)), str(path))
            rendered = source_from_docx_bytes(path.read_bytes())
            text = '\n'.join(b['text'] for b in rendered['blocks'])
            for block in source['blocks']:
                self.assertIn(block['text'], text)
            doc = Document(path)
            nested = doc.tables[0].cell(1, 1).tables[0]
            self.assertLessEqual(sum(c.width for c in nested.columns), doc.tables[0].cell(1, 1).width)

    def test_full_document_context_and_schema_contract(self):
        source = source_from_text('ชื่อหลักสูตร\nวัตถุประสงค์\nเข้าใจทีม\nรายละเอียดท้ายเอกสาร')
        prompt = build_prompt(source)
        self.assertIn('รายละเอียดท้ายเอกสาร', prompt)
        self.assertIn('STRUCTURE ANALYST', prompt)
        self.assertIn('paragraph CAN be split', prompt)

    def test_split_one_paragraph_into_nested_heading_and_description(self):
        text = 'Mindset: ฝึกตั้งคำถามเพื่อเข้าใจทีม'
        source = source_from_text(text)
        data = resolve_plan(source, plan([
            node('h', 'heading', 'B000001', text, end=8),
            node('p', 'paragraph', 'B000001', text, parent='h', start=8)]))
        self.assertFalse(data['_audit']['review_required'])
        self.assertEqual(data['_reference_document']['roots'][0]['children'][0]['kind'], 'paragraph')

    def test_thai_and_emoji_spans_use_unicode_codepoints(self):
        text = 'ทีม 🧠 เรียนรู้'
        source = source_from_text(text)
        data = resolve_plan(source, plan([node('p', 'paragraph', 'B000001', text)]))
        self.assertTrue(data['_audit']['all_source_content_retained'])

    def test_duplicate_span_is_rejected(self):
        source = source_from_text('duplicate')
        with self.assertRaisesRegex(JobProcessingError, 'duplicated'):
            resolve_plan(source, plan([node('a', 'paragraph', 'B000001', 'duplicate'), node('b', 'paragraph', 'B000001', 'duplicate')]))

    def test_identical_words_in_different_source_locations_are_retained(self):
        source = source_from_text('Workshop\nWorkshop')
        self.assertFalse(resolve_plan(source, identity_plan(source))['_audit']['review_required'])

    def test_missing_content_is_recovered_and_requires_review(self):
        source = source_from_text('หลักสูตร\nประโยคท้ายที่ห้ามหาย')
        data = resolve_plan(source, plan([node('t', 'title', 'B000001', 'หลักสูตร')]))
        self.assertTrue(data['_audit']['review_required'])
        self.assertEqual(data['_reference_document']['roots'][-1]['children'][0]['refs'][0]['block_id'], 'B000002')

    def test_invalid_or_invented_reference_fails(self):
        source = source_from_text('abc')
        for ref in ({'block_id': 'unknown', 'start': 0, 'end': 3}, {'block_id': 'B000001', 'start': 0, 'end': 99}, {'block_id': 'B000001', 'start': False, 'end': 3}):
            with self.subTest(ref=ref), self.assertRaises(JobProcessingError):
                item = node('a', 'paragraph'); item['refs'] = [ref]
                resolve_plan(source, plan([item]))

    def test_invalid_parent_and_time_cannot_silently_drop_content(self):
        source = source_from_text('09:00')
        for item in (node('a', 'paragraph', 'B000001', '09:00', parent='missing'), node('a', 'time', 'B000001', '09:00')):
            with self.assertRaises(JobProcessingError): resolve_plan(source, plan([item]))

    def test_original_section_order_is_enforced(self):
        source = source_from_text('วัตถุประสงค์\nประโยชน์')
        with self.assertRaisesRegex(JobProcessingError, 'reordered'):
            resolve_plan(source, plan([node('later', 'section', 'B000002', 'ประโยชน์'),
                                       node('earlier', 'section', 'B000001', 'วัตถุประสงค์')]))

    def test_arbitrary_model_authored_body_text_is_rejected(self):
        source = source_from_text('original')
        p = plan([node('a', 'paragraph', 'B000001', 'original')]); p['nodes'][0]['text'] = 'invented'
        with self.assertRaises(JobProcessingError): resolve_plan(source, p)

    def test_edit_proposals_do_not_change_source_by_default(self):
        source = source_from_text('teem')
        edits = [{'block_id': 'B000001', 'start': 0, 'end': 4, 'replacement': 'team', 'reason': 'spelling'}]
        data = resolve_plan(source, plan([node('a', 'paragraph', 'B000001', 'teem')], edits))
        self.assertFalse(data['_audit']['edits'][0]['applied'])
        self.assertEqual(data['_reference_document']['source']['blocks'][0]['text'], 'teem')
        self.assertTrue(data['_audit']['review_required'])

    def test_overlapping_edits_rejected(self):
        source = source_from_text('teem')
        edit = {'block_id': 'B000001', 'start': 0, 'end': 4, 'replacement': 'team', 'reason': 'spelling'}
        with self.assertRaisesRegex(JobProcessingError, 'overlap'):
            resolve_plan(source, plan([node('a', 'paragraph', 'B000001', 'teem')], [edit, edit]))

    def test_explicit_edit_application_is_reported_and_keeps_original_in_audit(self):
        source = source_from_text('teem')
        edits = [{'block_id': 'B000001', 'start': 0, 'end': 4, 'replacement': 'team', 'reason': 'spelling'}]
        data = resolve_plan(source, plan([node('a', 'paragraph', 'B000001', 'teem')], edits), apply_edits=True)
        self.assertTrue(data['_audit']['edits'][0]['applied'])
        with tempfile.TemporaryDirectory(dir=ROOT) as tmp:
            path = Path(tmp) / 'edited.docx'; generate_doc(data, str(path))
            paragraphs = [p.text for p in Document(path).paragraphs]
            self.assertEqual(paragraphs[0], 'team')
            self.assertTrue(any('teem → team' in p for p in paragraphs))

    def test_numbered_lists_from_word_styles_keep_original_labels(self):
        doc = Document(); doc.add_paragraph('First', 'List Number'); doc.add_paragraph('Second', 'List Number')
        source = source_from_docx_bytes(doc_bytes(doc))
        self.assertEqual([b['list_label'] for b in source['blocks']], ['1.', '2.'])
        with tempfile.TemporaryDirectory(dir=ROOT) as tmp:
            path = Path(tmp) / 'numbered.docx'; generate_doc(resolve_plan(source, identity_plan(source)), str(path))
            self.assertEqual([p.text for p in Document(path).paragraphs], ['1. First', '2. Second'])

    def test_audit_survives_job_store_restart_and_expires_with_result(self):
        with tempfile.TemporaryDirectory(dir=ROOT) as tmp:
            store = JobStore(tmp)
            def render(text, key, path):
                doc = Document(); doc.add_paragraph(text); doc.save(path)
                return 'Test.docx', 'mock', {'review_required': True, 'warnings': ['Review']}
            job = store.submit('audit-persist', 'source', 'dummy', render); store.close()
            store = JobStore(tmp)
            self.assertTrue(store.get(job['id'])['review_required'])
            audit_path = store.result_path(job['id']).with_suffix('.audit.json')
            self.assertTrue(audit_path.is_file())
            store.cleanup(retention_seconds=-1)
            self.assertFalse(audit_path.exists()); store.close()

    def test_docx_reader_preserves_runs_tables_and_horizontal_vertical_merges(self):
        doc = Document(); p = doc.add_paragraph('ไทย '); p.add_run('สำคัญ 🧠').bold = True
        table = doc.add_table(rows=3, cols=3)
        table.cell(0, 0).merge(table.cell(0, 2)).text = 'หัวตารางรวม'
        table.cell(1, 0).merge(table.cell(2, 0)).text = 'เซลล์แนวตั้ง'
        table.cell(1, 1).text = 'A'; table.cell(1, 2).text = 'B'
        table.cell(2, 1).text = 'C'; table.cell(2, 2).text = 'D'
        source = source_from_docx_bytes(doc_bytes(doc)); validate_source(source)
        self.assertEqual(source['blocks'][0]['bold'], [{'start': 4, 'end': 11}])
        self.assertEqual(source['tables'][0]['rows'][0]['cells'][0]['col_span'], 3)
        self.assertEqual(source['tables'][0]['rows'][1]['cells'][0]['row_span'], 2)
        with tempfile.TemporaryDirectory(dir=ROOT) as tmp:
            output = Path(tmp) / 'merged.docx'; generate_doc(resolve_plan(source, identity_plan(source)), str(output))
            result = Document(output)
            self.assertEqual(result.tables[0].cell(0, 0).text.strip(), 'หัวตารางรวม')
            self.assertEqual(result.tables[0].cell(1, 0).text.strip(), 'เซลล์แนวตั้ง')
            self.assertTrue(any(r.bold and r.text == 'สำคัญ 🧠' for r in result.paragraphs[0].runs))

    def test_source_table_cannot_be_flattened_or_duplicated(self):
        doc = Document(); doc.add_table(rows=1, cols=1).cell(0, 0).text = 'content'
        source = source_from_docx_bytes(doc_bytes(doc))
        with self.assertRaisesRegex(JobProcessingError, 'flattened'):
            resolve_plan(source, plan([node('a', 'paragraph', source['blocks'][0]['id'], 'content')]))
        with self.assertRaises(JobProcessingError):
            resolve_plan(source, plan([node('a', 'table', table='T000001'), node('b', 'table', table='T000001')]))

    def test_omitted_table_is_retained_for_review(self):
        doc = Document(); doc.add_paragraph('Course'); doc.add_table(rows=1, cols=1).cell(0, 0).text = 'Keep me'
        source = source_from_docx_bytes(doc_bytes(doc))
        data = resolve_plan(source, plan([node('a', 'title', 'B000001', 'Course')]))
        self.assertEqual(data['_audit']['missing_from_ai_plan'], [{'table_id': 'T000001'}])

    def test_renderer_handles_semantic_modules_time_workshops_and_nested_lists(self):
        texts = ['พัฒนาทีม', 'สิ่งที่ผู้เรียนจะทำได้', 'เมื่อจบการอบรม', 'เข้าใจทีม', 'Agenda',
                 'Module 1: ทีม', '09.00–10.30', 'Workshop: ทดลอง', 'อภิปราย', 'สรุปผล']
        source = source_from_text('\n'.join(texts))
        spec = [('t', 'title', '', ''), ('s1', 'section', '', 'objectives'), ('intro', 'paragraph', 's1', ''),
                ('b1', 'bullet', 's1', ''), ('s2', 'section', '', 'agenda'), ('m', 'module', 's2', ''),
                ('tm', 'time', 'm', ''), ('w', 'heading', 'm', ''), ('b2', 'bullet', 'w', ''), ('b3', 'bullet', 'b2', '')]
        nodes = [node(nid, kind, f'B{i:06d}', texts[i - 1], parent, category=cat) for i, (nid, kind, parent, cat) in enumerate(spec, 1)]
        with tempfile.TemporaryDirectory(dir=ROOT) as tmp:
            path = Path(tmp) / 'semantic.docx'; generate_doc(resolve_plan(source, plan(nodes)), str(path))
            doc = Document(path)
            self.assertEqual(len(doc.tables), 1)
            self.assertIn('09.00–10.30', doc.tables[0].cell(1, 0).text)
            self.assertIn('Workshop: ทดลอง', doc.tables[0].cell(1, 1).text)
            nested = next(p for p in doc.tables[0].cell(1, 1).paragraphs if p.text == 'สรุปผล')
            self.assertEqual(nested._p.find('./' + qn('w:pPr') + '/' + qn('w:numPr') + '/' + qn('w:ilvl')).get(qn('w:val')), '2')
            intro = next(p for p in doc.paragraphs if p.text == 'เมื่อจบการอบรม')
            self.assertIsNone(intro._p.find('.//' + qn('w:numPr')))

    def test_real_workspace_documents_roundtrip_all_extracted_body_text(self):
        files = list(ROOT.glob('B Tools_*.docx'))
        self.assertGreaterEqual(len(files), 8)
        with tempfile.TemporaryDirectory(dir=ROOT) as tmp:
            for index, file in enumerate(files):
                with self.subTest(file=file.name):
                    source = source_from_docx_bytes(file.read_bytes()); validate_source(source)
                    data = resolve_plan(source, identity_plan(source)); path = Path(tmp) / f'{index}.docx'
                    generate_doc(data, str(path)); result = source_from_docx_bytes(path.read_bytes())
                    rendered = '\n'.join(b['text'] for b in result['blocks'])
                    for b in source['blocks']:
                        self.assertIn(b['text'], rendered)


class IntegrationTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(dir=ROOT)
        self.store = JobStore(self.tmp.name); self.previous = api._job_store; api._job_store = self.store
        self.env = patch.dict(os.environ, {'BTOOLS_API_TOKEN': 'test-token', 'GEMINI_API_KEY': 'dummy', 'BTOOLS_PARSER_MODE': 'references'})
        self.env.start(); self.client = TestClient(api.app)
        self.headers = {'X-BTools-Token': 'test-token'}
        self.network = patch.object(api.requests, 'post', side_effect=AssertionError('Live network forbidden'))
        self.network.start()

    def tearDown(self):
        self.store.close(); api._job_store = self.previous
        self.client.close(); self.network.stop(); self.env.stop(); self.tmp.cleanup()

    def wait(self, jid):
        until = time.monotonic() + 5
        while time.monotonic() < until:
            job = self.store.get(jid)
            if job['status'] in ('failed', 'succeeded'): return job
            time.sleep(.01)
        self.fail('Job timed out')

    def test_docx_upload_reference_provider_audit_and_download_end_to_end(self):
        doc = Document(); doc.add_paragraph('Course'); doc.add_paragraph('Keep final paragraph')
        source = source_from_docx_bytes(doc_bytes(doc))
        incomplete = plan([node('t', 'title', 'B000001', 'Course')])
        class Response:
            status_code = 200
            def json(self): return {'candidates': [{'finishReason': 'STOP', 'content': {'parts': [{'text': json.dumps(incomplete)}]}}], 'usageMetadata': {'promptTokenCount': 123}}
        with patch.object(api.requests, 'post', return_value=Response()) as post:
            response = self.client.post('/api/jobs', headers=self.headers, json={'request_id': 'upload', 'source_docx_base64': base64.b64encode(doc_bytes(doc)).decode()})
            self.assertEqual(response.status_code, 202)
            jid = response.json()['id']; job = self.wait(jid)
            self.assertEqual(job['status'], 'succeeded'); self.assertTrue(job['review_required']); self.assertTrue(job['audit_available'])
            audit = self.client.get(f'/api/jobs/{jid}/audit', headers=self.headers).json()
            self.assertEqual(audit['token_usage']['promptTokenCount'], 123)
            self.assertTrue(audit['missing_from_ai_plan'])
            self.assertEqual(self.client.get(f'/api/jobs/{jid}/audit').status_code, 401)
            output = Document(BytesIO(self.client.get(f'/api/jobs/{jid}/result', headers=self.headers).content))
            self.assertIn('Keep final paragraph', [p.text for p in output.paragraphs])
            self.assertIn('responseJsonSchema', post.call_args.kwargs['json']['generationConfig'])
            self.assertNotIn('dummy', post.call_args.args[0])

    def test_changed_formatting_with_same_request_id_is_a_conflict(self):
        source = source_from_text('Course')
        with patch.object(api, 'render_background_job', return_value=('ignored.docx', 'mock')):
            body = {'raw_text': 'Course', 'source_document': source, 'request_id': 'formatting'}
            self.assertEqual(self.client.post('/api/jobs', json=body, headers=self.headers).status_code, 202)
            changed = copy.deepcopy(body); changed['source_document']['blocks'][0]['bold'] = [{'start': 0, 'end': 6}]
            self.assertEqual(self.client.post('/api/jobs', json=changed, headers=self.headers).status_code, 409)

    def test_bad_docx_and_mismatching_structured_text_rejected(self):
        for body in ({'source_docx_base64': 'invalid!'}, {'source_docx_base64': base64.b64encode(b'not zip').decode()},
                     {'raw_text': 'Different', 'source_document': source_from_text('Course')}):
            with self.subTest(body=body):
                self.assertEqual(self.client.post('/api/jobs', json={**body, 'request_id': 'invalid'}, headers=self.headers).status_code, 400)

    def test_emergency_legacy_switch_calls_original_parser(self):
        with patch.dict(os.environ, {'BTOOLS_PARSER_MODE': 'legacy'}), patch.object(api, 'call_gemini_api_legacy', return_value={'legacy': True}) as legacy:
            self.assertEqual(api.call_gemini_api('source', 'dummy'), {'legacy': True})
            legacy.assert_called_once_with('source', 'dummy')


if __name__ == '__main__':
    unittest.main(verbosity=2)
