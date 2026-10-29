"""Recovery protocol and hard worker deadlines. No real provider calls."""
import copy
import json
import os
from pathlib import Path
import subprocess
import sys
import tempfile
import time
import unittest
from unittest.mock import patch, Mock

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / '.test_dependencies'))
sys.path.insert(0, str(ROOT / 'btools_webapp'))
from fastapi.testclient import TestClient
from docx import Document
import app as api
from checkpoints import make_checkpoint, restore_checkpoint
from job_store import JobStore, JobProcessingError
from job_process import run_isolated


class FreeRecoveryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='test-recovery-', dir=ROOT)
        self.env = patch.dict(os.environ, {'BTOOLS_API_TOKEN': 'recovery-test-key', 'GEMINI_API_KEY': ''})
        self.env.start()
        self.previous = api._job_store
        self.store = api._job_store = JobStore(self.temp.name)
        self.client = TestClient(api.app)
        self.headers = {'X-BTools-Token': 'recovery-test-key'}
        self.data = {'course_title_en': 'Recovery', 'objectives': ['Keep original'], '_ai_model_used': 'mock'}
        self.network = patch.object(api.requests, 'post', side_effect=AssertionError('Live provider forbidden'))
        self.network.start()

    def tearDown(self):
        self.store.close(); self.client.close()
        api._job_store = self.previous
        self.network.stop(); self.env.stop(); self.temp.cleanup()

    def terminal(self, job_id):
        for _ in range(500):
            job = self.store.get(job_id)
            if job['status'] not in ('queued', 'processing'):
                return job
            time.sleep(.01)
        self.fail('Job did not finish')

    def test_checkpoint_then_storage_loss_then_render_without_ai(self):
        with patch.object(api, 'call_gemini_api', return_value=self.data) as ai:
            response = self.client.post('/api/jobs', headers=self.headers,
                json={'request_id': 'analysis', 'raw_text': 'original', 'api_key': 'dummy', 'checkpoint_only': True})
            self.assertEqual(response.status_code, 202)
            jid = response.json()['id']
            self.assertEqual(self.terminal(jid)['status'], 'checkpoint_ready')
            checkpoint = self.client.get(f'/api/jobs/{jid}/checkpoint', headers=self.headers).json()
            self.assertEqual(ai.call_count, 1)
        self.assertEqual(self.client.get(f'/api/jobs/{jid}/checkpoint').status_code, 401)
        self.assertEqual(self.client.get(f'/api/jobs/{jid}/result', headers=self.headers).status_code, 409)
        self.store.close()
        self.store = api._job_store = JobStore(Path(self.temp.name) / 'fresh-server')
        with patch.object(api, 'call_gemini_api', side_effect=AssertionError('Must reuse checkpoint')):
            response = self.client.post('/api/jobs', headers=self.headers,
                json={'request_id': 'render', 'raw_text': 'original', 'resume_checkpoint': checkpoint})
            self.assertEqual(response.status_code, 202)
            jid = response.json()['id']
            self.assertEqual(self.terminal(jid)['status'], 'succeeded')
        self.assertIn('Keep original', [p.text for p in Document(self.store.result_path(jid)).paragraphs])

    def test_tampered_or_wrong_source_checkpoint_rejected(self):
        cp = make_checkpoint('original', None, self.data)
        cases = [(cp, 'changed')]
        tampered = copy.deepcopy(cp); tampered['body']['data']['course_title_en'] = 'tampered'
        cases.extend([(tampered, 'original'), ({}, 'original')])
        for value, text in cases:
            response = self.client.post('/api/jobs', headers=self.headers,
                json={'request_id': 'invalid', 'raw_text': text, 'resume_checkpoint': value})
            self.assertEqual(response.status_code, 400)

    def test_same_identity_cannot_switch_phase(self):
        with patch.object(api, 'call_gemini_api', return_value=self.data):
            body = {'request_id': 'same', 'raw_text': 'original', 'api_key': 'dummy', 'checkpoint_only': True}
            first = self.client.post('/api/jobs', headers=self.headers, json=body)
            self.terminal(first.json()['id'])
            self.assertEqual(self.client.post('/api/jobs', headers=self.headers, json=body).json()['id'], first.json()['id'])
            body['checkpoint_only'] = False
            self.assertEqual(self.client.post('/api/jobs', headers=self.headers, json=body).status_code, 409)

    def test_real_isolated_process_restores_checkpoint(self):
        cp = make_checkpoint('original', None, self.data)
        path = str(Path(self.temp.name) / 'isolated.docx')
        result = run_isolated('original', '', path, None, {'resume_checkpoint': cp}, timeout=20)
        self.assertEqual(result[1], 'mock')
        self.assertIn('Keep original', [p.text for p in Document(path).paragraphs])

    def test_timeout_kills_and_reaps_before_reporting_failure(self):
        proc = Mock()
        proc.communicate.side_effect = [subprocess.TimeoutExpired('worker', .01), ('', '')]
        with patch('job_process.subprocess.Popen', return_value=proc):
            with self.assertRaises(JobProcessingError) as caught:
                run_isolated('text', 'secret', 'unused', None, {}, .01)
        self.assertEqual(caught.exception.code, 'worker_timeout')
        self.assertEqual([c[0] for c in proc.mock_calls], ['communicate', 'kill', 'communicate'])

    def test_timeout_frees_queue_for_next_job(self):
        from docx import Document
        def timeout(*args):
            raise JobProcessingError('Stopped', code='worker_timeout')
        self.store.isolated = True
        with patch('job_process.run_isolated', side_effect=timeout):
            job = self.store.submit('timeout', 'text', '', api.render_background_job)
            self.assertEqual(self.terminal(job['id'])['diagnostics']['code'], 'worker_timeout')
        self.store.isolated = False
        def render(text, keys, path):
            Document().save(path); return 'OK.docx', 'mock'
        next_job = self.store.submit('next', 'text', '', render)
        self.assertEqual(self.terminal(next_job['id'])['status'], 'succeeded')

    def test_quota_failure_is_distinct_from_configuration(self):
        with patch.dict(os.environ, {'GEMINI_MODELS': 'test-model'}):
            for status, code in [(429, 'quota_exhausted'), (401, 'configuration_error'), (503, 'provider_unavailable')]:
                response = Mock(status_code=status, headers={})
                with patch.object(api.requests, 'post', return_value=response), patch.object(api.time, 'sleep'):
                    with self.assertRaises(JobProcessingError) as caught:
                        api.call_gemini_api('original', 'dummy')
                self.assertEqual(caught.exception.code, code)


if __name__ == '__main__':
    unittest.main(verbosity=2)
