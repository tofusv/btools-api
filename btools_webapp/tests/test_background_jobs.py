"""Run this file directly; other legacy test scripts call live Gemini APIs."""
import importlib
import os
from pathlib import Path
import sqlite3
import sys
import tempfile
import threading
import time
import unittest
import uuid
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT / ".test_dependencies"))
sys.path.insert(0, str(ROOT / "btools_webapp"))
from fastapi.testclient import TestClient
from docx import Document
from job_store import JobStore, JobConflict, QueueFull

api = importlib.import_module("app")


class BackgroundJobsTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix="test-jobs-", dir=ROOT)
        self.assertTrue(Path(self.temp.name).resolve().is_relative_to(ROOT))
        self.store = JobStore(self.temp.name)
        self.old_store = api._job_store
        api._job_store = self.store
        self.env = patch.dict(os.environ, {"BTOOLS_API_TOKEN": "private-test-token", "GEMINI_API_KEY": "dummy"})
        self.env.start()
        self.network = patch.object(api.requests, "post", side_effect=AssertionError("Real Gemini calls disabled"))
        self.network.start()
        self.client = TestClient(api.app)
        self.headers = {"X-BTools-Token": "private-test-token"}
        self.gate = threading.Event()

    def tearDown(self):
        self.gate.set()
        self.store.close()
        self.client.close()
        api._job_store = self.old_store
        self.network.stop()
        self.env.stop()
        self.temp.cleanup()

    def wait_terminal(self, job_id):
        deadline = time.monotonic() + 5
        while time.monotonic() < deadline:
            job = self.store.get(job_id)
            if job["status"] in ("succeeded", "failed"):
                return job
            time.sleep(0.01)
        self.fail("job did not finish")

    def slow_render(self, text, keys, path):
        if not self.gate.wait(5):
            raise RuntimeError("test timed out")
        doc = Document()
        doc.add_paragraph(text)
        doc.save(path)
        return "B Tools_Test.docx", "mock-model"

    def test_large_job_returns_before_rendering_then_downloads(self):
        with patch.object(api, "render_background_job", self.slow_render):
            before = time.monotonic()
            response = self.client.post("/api/jobs", headers=self.headers,
                json={"raw_text": "Large document " * 10000, "request_id": "large-1"})
            self.assertEqual(response.status_code, 202)
            self.assertLess(time.monotonic() - before, 1)
            job_id = response.json()["id"]
            self.assertIn(response.json()["status"], ("queued", "processing"))
            self.assertEqual(self.client.get(f"/api/jobs/{job_id}/result", headers=self.headers).status_code, 409)
            self.gate.set()
            self.assertEqual(self.wait_terminal(job_id)["status"], "succeeded")
            result = self.client.get(f"/api/jobs/{job_id}/result", headers=self.headers)
            self.assertEqual(result.status_code, 200)
            self.assertTrue(result.content.startswith(b"PK"))
            self.assertEqual(result.headers["x-ai-model-used"], "mock-model")

    def test_resubmission_after_lost_response_is_idempotent(self):
        calls = []
        def render(text, keys, path):
            calls.append(text)
            return self.slow_render(text, keys, path)
        with patch.object(api, "render_background_job", render):
            payload = {"raw_text": "same source", "request_id": "stable-1"}
            a = self.client.post("/api/jobs", json=payload, headers=self.headers)
            b = self.client.post("/api/jobs", json=payload, headers=self.headers)
            self.assertEqual(a.json()["id"], b.json()["id"])
            conflict = self.client.post("/api/jobs", headers=self.headers,
                json={"raw_text": "edited source", "request_id": "stable-1"})
            self.assertEqual(conflict.status_code, 409)
            self.gate.set()
            self.wait_terminal(a.json()["id"])
            self.assertEqual(calls, ["same source"])

    def test_auth_validation_and_queue_limit(self):
        payload = {"raw_text": "source", "request_id": "one"}
        self.assertEqual(self.client.post("/api/jobs", json=payload).status_code, 401)
        self.assertEqual(self.client.post("/api/jobs", json={**payload, "raw_text": " "}, headers=self.headers).status_code, 400)
        with patch.object(api, "render_background_job", self.slow_render):
            for request_id in ("one", "two"):
                self.assertEqual(self.client.post("/api/jobs", json={**payload, "request_id": request_id}, headers=self.headers).status_code, 202)
            full = self.client.post("/api/jobs", json={**payload, "request_id": "three"}, headers=self.headers)
            self.assertEqual(full.status_code, 429)

    def test_failure_does_not_disclose_keys_and_has_terminal_state(self):
        def fail(text, keys, path):
            raise RuntimeError("secret Gemini key: " + keys)
        job = self.store.submit("failure", "text", "SECRET_KEY", fail)
        result = self.wait_terminal(job["id"])
        self.assertEqual(result["status"], "failed")
        self.assertNotIn("SECRET_KEY", result["error"])

    def test_restart_marks_interrupted_jobs_failed(self):
        self.store.close()
        job_id = str(uuid.uuid4())
        with sqlite3.connect(Path(self.temp.name) / "jobs.sqlite3") as db:
            db.execute("INSERT INTO jobs (id,request_id,digest,status,created,updated) VALUES (?,?,?,'processing',?,?)",
                       (job_id, "interrupted", "digest", time.time(), time.time()))
        db.close()
        self.store = JobStore(self.temp.name)
        api._job_store = self.store
        self.assertEqual(self.store.get(job_id)["status"], "failed")
        self.assertIn("restarted", self.store.get(job_id)["error"])

    def test_success_persists_and_expired_jobs_are_removed(self):
        self.gate.set()
        job = self.store.submit("persist", "text", "dummy", self.slow_render)
        self.wait_terminal(job["id"])
        self.store.close()
        self.store = JobStore(self.temp.name)
        api._job_store = self.store
        self.assertEqual(self.store.get(job["id"])["status"], "succeeded")
        self.assertTrue(self.store.result_path(job["id"]).is_file())
        self.store.cleanup(retention_seconds=-1)
        self.assertIsNone(self.store.get(job["id"]))
        self.assertFalse(self.store.result_path(job["id"]).exists())

    def test_real_renderer_with_mock_ai_generates_docx(self):
        with patch.object(api, "call_gemini_api", return_value={"course_title_en": "Sample", "objectives": ["Learn"], "_ai_model_used": "mock"}):
            response = self.client.post("/api/jobs", headers=self.headers,
                json={"raw_text": "source", "request_id": "real-docx"})
            job = self.wait_terminal(response.json()["id"])
            self.assertEqual(job["status"], "succeeded")
            doc = Document(self.store.result_path(job["id"]))
            self.assertIn("Learn", [p.text for p in doc.paragraphs])

    def test_truncated_ai_output_is_reported_immediately(self):
        class Response:
            status_code = 200
            def json(self): return {"candidates": [{"finishReason": "MAX_TOKENS"}]}
        with patch.object(api.requests, "post", return_value=Response()):
            with self.assertRaisesRegex(api.JobProcessingError, "truncated"):
                api.call_gemini_api("source", "dummy")


if __name__ == "__main__":
    unittest.main(verbosity=2)
