"""Render the validated reference tree with B Tools typography and source runs."""
import re
from docx.shared import Inches, Pt
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


def render_reference_document(doc, payload, fmt):
    source = payload['source']
    blocks = {b['id']: b for b in source['blocks']}
    tables = {t['id']: t for t in source.get('tables', [])}
    edits = [e for e in payload['audit']['edits'] if e['applied']]
    # Own numbering definition: templates may use numId=1 for a numbered list.
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(e.get(qn('w:abstractNumId'))) for e in numbering.findall(qn('w:abstractNum'))]
    num_ids = [int(e.get(qn('w:numId'))) for e in numbering.findall(qn('w:num'))]
    abstract_id, num_id = max(abstract_ids + [-1]) + 1, max(num_ids + [0]) + 1
    abstract = OxmlElement('w:abstractNum')
    abstract.set(qn('w:abstractNumId'), str(abstract_id))
    for level in range(9):
        lvl = OxmlElement('w:lvl'); lvl.set(qn('w:ilvl'), str(level))
        for tag, value in [('start', '1'), ('numFmt', 'bullet'), ('lvlText', '•')]:
            item = OxmlElement('w:' + tag); item.set(qn('w:val'), value); lvl.append(item)
        abstract.append(lvl)
    numbering.append(abstract)
    num = OxmlElement('w:num'); num.set(qn('w:numId'), str(num_id))
    abstract_ref = OxmlElement('w:abstractNumId'); abstract_ref.set(qn('w:val'), str(abstract_id))
    num.append(abstract_ref); numbering.append(num)
    doc.styles['Normal'].font.name = fmt.STRICT_FONT_NAME
    doc.styles['Normal'].font.size = Pt(10)

    def run(p, text, bold=False, size=10):
        r = p.add_run(text); r.font.name = fmt.STRICT_FONT_NAME
        r.font.size = Pt(size); r.bold = bold
        fonts = r._r.get_or_add_rPr().rFonts
        fonts.set(qn('w:cs'), fmt.STRICT_FONT_NAME)
        fonts.set(qn('w:eastAsia'), fmt.STRICT_FONT_NAME)

    def add_refs(p, refs, heading=False, size=10, bullet=False):
        for i, ref in enumerate(refs):
            block = blocks[ref['block_id']]
            start, end = ref['start'], ref['end']
            if i and refs[i - 1]['block_id'] != ref['block_id']:
                run(p, ' ', heading, size)
            if start == 0 and block.get('list_label'):
                run(p, block['list_label'] + ' ', heading, size)
            if bullet and i == 0:
                prefix = re.match(r'^\s*[•\-]\s+', block['text'][start:end])
                if prefix: start += prefix.end()
            changes = [e for e in edits if e['block_id'] == ref['block_id'] and start <= e['start'] < e['end'] <= end]
            boundaries = {start, end}
            for span in block.get('bold', []):
                boundaries.update((max(start, min(end, span['start'])), max(start, min(end, span['end']))))
            for e in changes:
                boundaries.update((e['start'], e['end']))
            cursor = start
            while cursor < end:
                change = next((e for e in changes if e['start'] == cursor), None)
                bold = heading or any(s['start'] <= cursor < s['end'] for s in block.get('bold', []))
                if change:
                    run(p, change['replacement'], bold, size); cursor = change['end']
                else:
                    stop = min(n for n in boundaries if n > cursor)
                    run(p, block['text'][cursor:stop], bold, size); cursor = stop

    def paragraph(container, node, depth=0, in_cell=False, category=''):
        kind = node['kind']
        p = container.add_paragraph()
        p.paragraph_format.space_after = Pt(3 if in_cell else 4)
        p.paragraph_format.line_spacing = 1.35 if in_cell else 1.5
        bold = kind in ('title', 'subtitle', 'section', 'heading', 'module')
        size = 13 if kind in ('title', 'subtitle') else (12 if kind == 'section' else 10)
        if bold:
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_before = Pt(8 if kind == 'section' else 3)
        if kind == 'paragraph' and category == 'rationale' and not in_cell:
            p.paragraph_format.first_line_indent = Inches(0.5)
        if kind == 'bullet':
            text = ''.join(blocks[r['block_id']]['text'][r['start']:r['end']] for r in node['refs'])
            source_label = blocks[node['refs'][0]['block_id']].get('list_label', '') if node['refs'] else ''
            numbered = bool(source_label) or bool(re.match(r'^\s*(\d+[.)]|[A-Za-z][.)]|->|=>|➢|➔)\s', text))
            if not numbered:
                num_pr = OxmlElement('w:numPr')
                ilvl = OxmlElement('w:ilvl'); ilvl.set(qn('w:val'), str(min(depth, 8)))
                num_ref = OxmlElement('w:numId'); num_ref.set(qn('w:val'), str(num_id))
                num_pr.append(ilvl); num_pr.append(num_ref); p._p.get_or_add_pPr().append(num_pr)
            p.paragraph_format.left_indent = Inches((0.12 if in_cell else 0.5) + 0.18 * depth)
            p.paragraph_format.first_line_indent = Inches(-0.12 if in_cell else -0.25)
        if node.get('review_heading'):
            run(p, node['review_heading'], True, size)
        else:
            add_refs(p, node['refs'], bold, size, kind == 'bullet')
        return p

    def prepare_table(container, rows, cols, widths=None):
        # Both Document and _Cell take rows/cols; _Cell rejects width.
        table = container.add_table(rows=rows, cols=cols)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        available = (container.width.inches if container.width else 6.1) if container is not doc else 6.65
        fmt.set_table_col_widths(table, widths or [available / cols] * cols)
        for row in table.rows:
            for cell in row.cells:
                fmt.set_cell_margins(cell)
        return table

    def source_table(container, node):
        original = tables[node['table_id']]
        cols = max((c['column'] + c.get('col_span', 1) for row in original['rows'] for c in row['cells']), default=1)
        table = prepare_table(container, len(original['rows']), cols)
        for ri, row in enumerate(original['rows']):
            for entry in row['cells']:
                col, cs, rs = entry['column'], entry.get('col_span', 1), entry.get('row_span', 1)
                cell = table.cell(ri, col)
                if cs > 1 or rs > 1:
                    cell = cell.merge(table.cell(ri + rs - 1, col + cs - 1))
                for p in list(cell.paragraphs): p._p.getparent().remove(p._p)
                for bid in entry['block_ids']:
                    block = blocks[bid]
                    if not block['text']: continue
                    paragraph(cell, {'kind': 'bullet' if block.get('kind') == 'bullet' else 'paragraph',
                                     'refs': [{'block_id': bid, 'start': 0, 'end': len(block['text'])}]},
                              block.get('level', 0), True)
                if not cell.paragraphs: cell.add_paragraph()
        fmt.apply_table_borders(table)
        if container is doc: fmt.add_blank_line(doc)

    def modules(container, nodes, depth, category):
        has_time = any(any(c['kind'] == 'time' for c in n['children']) for n in nodes)
        table = prepare_table(container, len(nodes) + 1, 2 if has_time else 1,
                              [1.55, 5.10] if has_time and container is doc else None)
        headers = ['เวลา', 'รายละเอียด'] if has_time else ['รายละเอียดเนื้อหาการอบรม']
        for cell, title in zip(table.rows[0].cells, headers):
            cell.paragraphs[0].paragraph_format.space_after = Pt(3)
            run(cell.paragraphs[0], title, True)
        for index, node in enumerate(nodes, 1):
            content = table.cell(index, 1 if has_time else 0)
            if has_time:
                time_cell = table.cell(index, 0)
                for p in list(time_cell.paragraphs): p._p.getparent().remove(p._p)
                for child in node['children']:
                    if child['kind'] == 'time':
                        paragraph(time_cell, child, in_cell=True).alignment = WD_ALIGN_PARAGRAPH.CENTER
                if not time_cell.paragraphs: time_cell.add_paragraph()
            for p in list(content.paragraphs): p._p.getparent().remove(p._p)
            if node['refs']: paragraph(content, node, in_cell=True)
            sequence(content, [c for c in node['children'] if c['kind'] != 'time'], depth, True, category)
            if not content.paragraphs: content.add_paragraph()
        fmt.apply_table_borders(table)
        if container is doc: fmt.add_blank_line(doc)

    def sequence(container, nodes, depth=0, in_cell=False, category=''):
        index = 0
        while index < len(nodes):
            node = nodes[index]
            if node['kind'] == 'module':
                end = index + 1
                while end < len(nodes) and nodes[end]['kind'] == 'module': end += 1
                modules(container, nodes[index:end], depth, category)
                index = end; continue
            if node['kind'] == 'table':
                source_table(container, node)
            elif node['refs'] or node.get('review_heading'):
                paragraph(container, node, depth, in_cell, category)
            if node['children']:
                sequence(container, node['children'], depth + (1 if node['kind'] in ('bullet', 'heading') else 0),
                         in_cell, node.get('category') or category)
            index += 1
    sequence(doc, payload['roots'])
    audit = payload['audit']
    if audit['review_required']:
        p = doc.add_paragraph(); p.paragraph_format.page_break_before = True
        run(p, 'รายการที่ต้องตรวจสอบก่อนนำเอกสารไปใช้', True, 12)
        run(doc.add_paragraph(), 'ส่วนนี้เป็นรายงานตรวจสอบ แยกจากเนื้อหาหลักสูตร')
        if audit['missing_from_ai_plan']:
            run(doc.add_paragraph(), 'มีข้อความหรือตารางที่ AI ยังจัดโครงสร้างไม่ครบ ระบบเก็บไว้ในหมวดเนื้อหาต้นฉบับที่รอตรวจสอบแล้ว')
        if audit['uncertain_nodes']:
            run(doc.add_paragraph(), 'AI ระบุหัวข้อที่ต้องตรวจโครงสร้าง: ' + ', '.join(audit['uncertain_nodes']))
        for warning in audit['warnings']:
            run(doc.add_paragraph(), warning)
        for edit in audit['edits']:
            label = 'แก้ไขแล้ว' if edit['applied'] else 'ข้อเสนอแก้คำ ยังไม่ได้แก้ต้นฉบับ'
            run(doc.add_paragraph(), label + ': ' + edit['original'] + ' → ' + edit['replacement'])
            run(doc.add_paragraph(), 'เหตุผลที่ AI เสนอ: ' + edit['reason'])
